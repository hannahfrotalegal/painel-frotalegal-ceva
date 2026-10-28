#!/usr/bin/env python3
"""
Acumula as inspeções do prontuário (aba "Inspeções últimos 60 dias") em
dados/historico_inspecoes.csv, para o painel enxergar além da janela de 60 dias.

Regra de mesclagem (sem risco de duplicar):
  histórico novo = linhas antigas com data ANTERIOR à 1ª data do arquivo novo
                   + todas as linhas do arquivo novo.

Uso:  python scripts/atualizar_historico.py [xlsx] [csv]
Padrão: dados/prontuario-equipamentos.xlsx  dados/historico_inspecoes.csv
Roda sozinho via GitHub Actions a cada upload do prontuário.
"""
import csv
import datetime as dt
import os
import sys

import openpyxl

ABA = "Inspeções últimos 60 dias"
CAMPOS = ["data", "identificador", "marca_modelo", "transportador", "tipo_inspecao",
          "checklist", "responsavel", "perfil", "unidade"]
ORIGEM = {  # coluna do prontuário -> campo do CSV
    "Identificador": "identificador", "Marca / Modelo": "marca_modelo",
    "Transportador": "transportador", "Tipo de Inspeção": "tipo_inspecao",
    "Checklist": "checklist", "Responsável": "responsavel",
    "Perfil Responsável": "perfil", "Unidade": "unidade",
}


def para_data(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def ler_prontuario(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if ABA not in wb.sheetnames:
        raise SystemExit(f"Aba '{ABA}' não encontrada em {caminho}")
    linhas = list(wb[ABA].iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(linhas[:10]) if r and "Identificador" in r), None)
    if hi is None:
        raise SystemExit("Cabeçalho com 'Identificador' não encontrado.")
    cab = [str(c or "").strip() for c in linhas[hi]]
    idx = {nome: cab.index(nome) for nome in list(ORIGEM) + ["Date de Execução"] if nome in cab}
    out = []
    for r in linhas[hi + 1:]:
        if not r or not r[idx["Identificador"]]:
            continue
        d = para_data(r[idx["Date de Execução"]])
        if not d:
            continue
        reg = {"data": d}
        for nome, campo in ORIGEM.items():
            v = r[idx[nome]] if nome in idx else ""
            reg[campo] = "" if v is None else str(v).strip()
        out.append(reg)
    return out


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "dados/prontuario-equipamentos.xlsx"
    csvp = sys.argv[2] if len(sys.argv) > 2 else "dados/historico_inspecoes.csv"
    novos = ler_prontuario(xlsx)
    if not novos:
        print("Arquivo sem inspeções — histórico mantido.")
        return
    corte = min(r["data"] for r in novos)
    antigos = []
    if os.path.exists(csvp):
        with open(csvp, encoding="utf-8-sig", newline="") as f:
            antigos = [r for r in csv.DictReader(f) if r.get("data", "") < corte]
    todos = sorted(antigos + novos, key=lambda r: (r["data"], r["identificador"]), reverse=True)
    with open(csvp, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CAMPOS, extrasaction="ignore")
        w.writeheader()
        w.writerows(todos)
    print(f"Histórico: {len(antigos)} linhas antigas (< {corte}) + {len(novos)} do arquivo = {len(todos)}")


if __name__ == "__main__":
    main()
